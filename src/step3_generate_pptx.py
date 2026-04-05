"""
Step 3: intermediates/ のMDファイルと画像を使ってPowerPointを生成する
使用ライブラリ: python-pptx

スライド構成:
  Slide 1: 表紙（外観写真＋物件名・キャッチコピー・価格・利回り）
  Slide 2: 内観（写真4枚均等グリッド・テキストなし）
  Slide 3: コンセプト（写真2枚＋生成コンセプト文＋エリア特性）
  Slide 4: FLOOR PLAN（左メイン図面＋右列各階＋下部建物スペック）
  Slide 5: RENT ROLL（タイプ別集計＋階ごと部屋グリッド＋合計）
  Slide 6: 物件立地（地図2枚縦積み＋物件概要・法令制限・アクセス）

共通要素: 右上に英語物件名 + ページ番号
"""

import re
from pathlib import Path
import yaml
from pptx import Presentation
from pptx.util import Mm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from image_utils import resize_for_pptx

SCRIPT_DIR = Path(__file__).parent
REPO_ROOT = SCRIPT_DIR.parent

# --- カラーパレット ---
COLOR_BLACK      = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GRAY       = RGBColor(0x88, 0x88, 0x88)
COLOR_LIGHT_GRAY  = RGBColor(0xD8, 0xD8, 0xD8)
COLOR_ULTRA_LIGHT = RGBColor(0xEC, 0xEC, 0xEC)  # セル内列区切り用（さらに薄い）
COLOR_ACCENT     = RGBColor(0xC8, 0xAA, 0x7A)  # ゴールド系アクセント
COLOR_BG         = RGBColor(0xF8, 0xF6, 0xF3)  # オフホワイト背景


# =========================================================
# ユーティリティ
# =========================================================

def load_config() -> dict:
    with open(REPO_ROOT / "config.yaml", "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def mm(value: float) -> Emu:
    return Mm(value)


def parse_md_value(md_text: str, key: str) -> str:
    """key: value 形式の値をMDから抽出する"""
    pattern = rf"^{re.escape(key)}:\s*(.+)$"
    m = re.search(pattern, md_text, re.MULTILINE)
    if m:
        return m.group(1).strip()
    return ""


def parse_md_multiline(md_text: str, key: str) -> str:
    """key: |\\n  行1\\n  行2 形式を抽出する"""
    pattern = rf"^{re.escape(key)}:\s*\|\s*\n((?:  .+\n?)*)"
    m = re.search(pattern, md_text, re.MULTILINE)
    if m:
        lines = m.group(1).split("\n")
        return "\n".join(l.strip() for l in lines if l.strip())
    return parse_md_value(md_text, key)


def parse_section_table(md_text: str, section_name: str) -> list[list[str]]:
    """指定セクション（## section_name）内の最初のテーブルのみ取得"""
    in_section = False
    rows = []
    for line in md_text.split("\n"):
        if line.startswith("##"):
            in_section = section_name in line
            continue
        if not in_section:
            continue
        if "|" in line:
            cells = [c.strip() for c in line.split("|") if c.strip()]
            if all(set(c) <= set("-:") for c in cells):
                continue  # セパレータ行スキップ
            rows.append(cells)
    return rows


def group_rooms_by_floor(table_rows: list[list[str]]) -> dict[int, list[list[str]]]:
    """部屋番号の先頭桁から階数を判定し5F→1Fの降順でグルーピング。
    ヘッダー行はスキップ。数値変換できない部屋番号（D棟等）もスキップ。"""
    floors: dict[int, list[list[str]]] = {}
    for row in table_rows[1:]:  # ヘッダー行スキップ
        if not row:
            continue
        try:
            floor = int(row[0].strip()) // 100
        except ValueError:
            continue
        floors.setdefault(floor, []).append(row)
    return dict(sorted(floors.items(), reverse=True))  # 5F→1F


def derive_type_summary(table_rows: list[list[str]]) -> list[tuple[str, str, int]]:
    """room tableからタイプ別集計を導出: [(タイプ名, 間取り, 戸数), ...]
    カラム順: [0]部屋番号 [1]用途 [2]占有面積 [3]タイプ [4]間取り [5]賃料 [6]坪単価 [7]状況"""
    type_map: dict[str, dict] = {}
    for row in table_rows[1:]:
        if len(row) < 5:
            continue
        try:
            int(row[0].strip())
        except ValueError:
            continue
        madori    = row[4].strip()
        type_name = row[3].strip()
        if type_name not in type_map:
            type_map[type_name] = {"madori": madori, "count": 0}
        type_map[type_name]["count"] += 1
    return [(t, v["madori"], v["count"]) for t, v in sorted(type_map.items())]


def derive_contract_status(table_rows: list[list[str]]) -> dict[str, int]:
    """room tableから契約状況サマリーを導出。
    カラム順: [0]部屋番号 [1]用途 [2]占有面積 [3]タイプ [4]間取り [5]賃料 [6]坪単価 [7]状況"""
    STATUS_COL = 7  # 状況カラム（0-based）
    counts = {"契約中": 0, "空室": 0, "退去予定": 0, "申込": 0}
    for row in table_rows[1:]:
        if len(row) <= STATUS_COL:
            continue  # 状況列がなければスキップ（全て0のまま）
        try:
            int(row[0].strip())
        except ValueError:
            continue
        status = row[STATUS_COL].strip()
        if status in counts:
            counts[status] += 1
    return counts


def read_md(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8")


def find_images(directory: Path, processed_dir: Path, page_key: str) -> list[Path]:
    """加工済み画像が存在すればそちらを、なければ元画像を返す"""
    processed = sorted(
        [p for p in processed_dir.glob(f"{page_key}_*")
         if p.suffix.lower() in (".jpg", ".jpeg", ".png")]
    )
    if processed:
        return processed
    if not directory.exists():
        return []
    return sorted(
        [p for p in directory.glob("*") if p.suffix.lower() in (".jpg", ".jpeg", ".png")]
    )


def find_image_by_prefix(directory: Path, processed_dir: Path, page_key: str, prefix: str) -> Path | None:
    """指定prefixに一致する加工済みまたは元画像を1枚返す（加工済み優先）"""
    exts = (".jpg", ".jpeg", ".png")
    processed = sorted([
        p for p in processed_dir.glob(f"{page_key}_{prefix}*")
        if p.suffix.lower() in exts
    ])
    if processed:
        return processed[0]
    if not directory.exists():
        return None
    matched = sorted([
        p for p in directory.glob(f"{prefix}*")
        if p.suffix.lower() in exts
    ])
    return matched[0] if matched else None


# =========================================================
# 描画ヘルパー
# =========================================================

def set_slide_background(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(
    slide, text: str,
    left, top, width, height,
    font_size: int = 11,
    bold: bool = False,
    color: RGBColor = COLOR_BLACK,
    align=PP_ALIGN.LEFT,
    font_name: str = "Noto Sans JP",
):
    txBox = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return txBox


def add_multiline_textbox(slide, lines: list[tuple], left, top, width, height,
                           align=PP_ALIGN.CENTER):
    """複数行テキストボックス。lines: [(text, font_size_pt, bold, RGBColor), ...]"""
    txBox = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = txBox.text_frame
    tf.word_wrap = False
    for i, (text, font_size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Noto Sans JP"
    return txBox


def add_image_safe(slide, image_path: Path, left, top, width, height):
    """object-fit: cover 相当。長辺をトリミングして枠に合わせる"""
    if not (image_path and image_path.exists()):
        return False
    resized = resize_for_pptx(image_path, max_side=2048)
    from PIL import Image as _Image
    img_w, img_h = _Image.open(resized).size
    target_w, target_h = int(width), int(height)
    img_aspect    = img_w / img_h
    target_aspect = target_w / target_h
    pic = slide.shapes.add_picture(str(resized), int(left), int(top), target_w, target_h)
    if img_aspect > target_aspect:
        visible_w = img_h * target_aspect
        crop_lr = (img_w - visible_w) / 2 / img_w
        pic.crop_left = crop_lr
        pic.crop_right = crop_lr
    else:
        visible_h = img_w / target_aspect
        crop_tb = (img_h - visible_h) / 2 / img_h
        pic.crop_top = crop_tb
        pic.crop_bottom = crop_tb
    pic.line.fill.background()  # 枠線なし
    return True


def add_image_fit(slide, image_path: Path, left, top, width, height):
    """object-fit: contain 相当。枠内に収まるようスケール（地図・図面向け）"""
    if not (image_path and image_path.exists()):
        return False
    resized = resize_for_pptx(image_path, max_side=2048)
    from PIL import Image as _Image
    img_w, img_h = _Image.open(resized).size
    target_w, target_h = int(width), int(height)
    img_aspect    = img_w / img_h
    target_aspect = target_w / target_h
    if img_aspect > target_aspect:
        actual_w  = target_w
        actual_h  = int(target_w / img_aspect)
        offset_top = (target_h - actual_h) // 2
        slide.shapes.add_picture(str(resized), int(left), int(top) + offset_top, actual_w, actual_h)
    else:
        actual_h   = target_h
        actual_w   = int(target_h * img_aspect)
        offset_left = (target_w - actual_w) // 2
        slide.shapes.add_picture(str(resized), int(left) + offset_left, int(top), actual_w, actual_h)
    return True


def add_header_common(slide, name_en: str, page_num: int, slide_width, margin_right):
    """右上に英語物件名とページ番号を追加（共通要素）"""
    add_text_box(
        slide, f"{name_en}  |  {page_num:02d}",
        left=slide_width - mm(80), top=mm(6), width=mm(75), height=mm(8),
        font_size=7, color=COLOR_GRAY, align=PP_ALIGN.RIGHT, font_name="Helvetica Neue",
    )


def add_separator_line(slide, left, top, width, color: RGBColor = COLOR_ACCENT):
    """水平セパレーターラインを追加"""
    line = slide.shapes.add_connector(1, int(left), int(top), int(left + width), int(top))
    line.line.color.rgb = color
    line.line.width = Pt(0.5)


def add_filled_rect(slide, left, top, width, height, fill_color: RGBColor):
    """塗りつぶし矩形を追加（テキストなし背景用）"""
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE
        int(left), int(top), int(width), int(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()  # 枠線なし
    return shape


def add_vertical_line(slide, left, top, height, color: RGBColor = COLOR_LIGHT_GRAY, width_pt: float = 0.5):
    """垂直セパレーターラインを追加"""
    line = slide.shapes.add_connector(1, int(left), int(top), int(left), int(top + height))
    line.line.color.rgb = color
    line.line.width = Pt(width_pt)




# =========================================================
# 各スライド生成関数
# =========================================================

def add_slide1_cover(prs, config, text_dir, img_dir, processed_dir):
    """スライド1: 表紙"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_BG)

    W  = prs.slide_width
    H  = prs.slide_height
    mr = mm(config["slide"]["margin_right_mm"])
    mt = mm(config["slide"]["margin_top_mm"])

    property_info = read_md(text_dir / "property_info.md")
    concept_text  = read_md(text_dir / "concept_text.md")

    name_ja    = parse_md_value(property_info, "name_ja")    or config["property"]["name_ja"]
    name_en    = parse_md_value(property_info, "name_en")    or config["property"]["name_en"]
    price      = parse_md_value(property_info, "price")
    yield_rate = parse_md_value(property_info, "yield_rate")
    catchcopy  = parse_md_value(concept_text,  "catchcopy")

    images = find_images(
        REPO_ROOT / config["paths"]["assets"] / "page1_image",
        processed_dir, "page1",
    )

    # 左: メイン外観写真（余白付き）
    main_img_w = W * 0.62
    img_pad_l = mm(10)
    img_pad_r = mm(6)   # 右セクションとの間隔
    img_pad_b = mm(10)
    if images:
        add_image_safe(
            slide, images[0],
            left=img_pad_l,
            top=mt,
            width=main_img_w - img_pad_l - img_pad_r,
            height=H - mt - img_pad_b,
        )

    # 右上: サブ外観写真（上端を左画像・他スライドと揃える）
    right_x   = main_img_w + mm(4)
    right_w   = W - right_x - mm(10)   # 右端にも余白
    sub_img_h = H * 0.44               # 高さ約44%
    if len(images) >= 2:
        add_image_safe(slide, images[1], right_x, mt, right_w, sub_img_h)

    # 右下: テキストエリア（サブ画像の下から）
    text_y = mt + sub_img_h + mm(7)
    text_w = right_w - mm(3)

    add_text_box(
        slide, name_ja,
        left=right_x, top=text_y, width=text_w, height=mm(11),
        font_size=12, bold=True, color=COLOR_BLACK, font_name="Noto Sans JP",
    )
    add_text_box(
        slide, name_en,
        left=right_x, top=text_y + mm(12), width=text_w, height=mm(7),
        font_size=7, color=COLOR_GRAY, font_name="Helvetica Neue",
    )
    add_separator_line(slide, right_x, text_y + mm(20), text_w)
    add_text_box(
        slide, catchcopy,
        left=right_x, top=text_y + mm(22), width=text_w, height=mm(18),
        font_size=7, color=COLOR_GRAY, font_name="Helvetica Neue",
    )
    if price:
        add_text_box(
            slide, price,
            left=right_x, top=text_y + mm(42), width=text_w, height=mm(9),
            font_size=12, bold=True, color=COLOR_ACCENT, font_name="Helvetica Neue",
        )
    if yield_rate:
        add_text_box(
            slide, f"想定表面利回り  {yield_rate}",
            left=right_x, top=text_y + mm(52), width=text_w, height=mm(7),
            font_size=8, color=COLOR_GRAY, font_name="Noto Sans JP",
        )

    add_header_common(slide, name_en, 1, W, mr)
    return slide


def add_slide2_interior(prs, config, text_dir, img_dir, processed_dir):
    """スライド2: 内観（写真4枚均等グリッド・余白あり）"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_BG)

    W  = prs.slide_width
    H  = prs.slide_height
    mr = mm(config["slide"]["margin_right_mm"])

    name_en = config["property"]["name_en"]
    images  = find_images(
        REPO_ROOT / config["paths"]["assets"] / "page2_image",
        processed_dir, "page2",
    )

    # 外周余白（ヘッダ表示のため上は広めに）
    pad_top    = mm(15)
    pad_bottom = mm(10)
    pad_left   = mm(12)
    pad_right  = mm(12)
    gap        = mm(3)   # 画像間の隙間

    area_w = W - pad_left - pad_right
    area_h = H - pad_top  - pad_bottom
    cell_w = (area_w - gap) / 2
    cell_h = (area_h - gap) / 2

    positions = [
        (pad_left,            pad_top,             cell_w, cell_h),
        (pad_left + cell_w + gap, pad_top,          cell_w, cell_h),
        (pad_left,            pad_top + cell_h + gap, cell_w, cell_h),
        (pad_left + cell_w + gap, pad_top + cell_h + gap, cell_w, cell_h),
    ]
    for i, (l, t, w, h) in enumerate(positions):
        if i < len(images):
            add_image_safe(slide, images[i], l, t, w, h)

    add_header_common(slide, name_en, 2, W, mr)
    return slide


def add_slide3_concept(prs, config, text_dir, img_dir, processed_dir):
    """スライド3: コンセプト（写真2枚＋生成コンセプト文＋エリア特性）"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_BG)

    W  = prs.slide_width
    H  = prs.slide_height
    ml = mm(config["slide"]["margin_left_mm"])
    mt = mm(config["slide"]["margin_top_mm"])
    mr = mm(config["slide"]["margin_right_mm"])

    name_en      = config["property"]["name_en"]
    concept_text = read_md(text_dir / "concept_text.md")
    location_info = read_md(text_dir / "location_info.md")

    concept_title    = parse_md_value(concept_text, "concept_title")
    concept_body     = parse_md_multiline(concept_text, "concept_body")
    features_raw     = re.findall(r"- feature\d+:\s*(.+)", concept_text)
    area_description = parse_md_multiline(location_info, "area_description") \
                       or parse_md_value(location_info, "area_description")

    images = find_images(
        REPO_ROOT / config["paths"]["assets"] / "page3_image",
        processed_dir, "page3",
    )

    usable_w = W - ml - mr
    gap      = mm(4)
    img_w    = (usable_w - gap) / 2
    img_h    = H * 0.55

    # 上段: コンセプト写真2枚
    if images:
        add_image_safe(slide, images[0], ml, mt - mm(5), img_w, img_h)
    if len(images) >= 2:
        add_image_safe(slide, images[1], ml + img_w + gap, mt - mm(5), img_w, img_h)

    # 下段テキストエリア
    text_top = mt + img_h - mm(2)

    add_text_box(
        slide, "D E S I G N  C O N C E P T",
        left=ml, top=text_top, width=mm(90), height=mm(8),
        font_size=7, bold=True, color=COLOR_ACCENT, font_name="Playfair Display",
    )
    add_separator_line(slide, ml, text_top + mm(9), usable_w)

    # 左: コンセプトタイトル + 本文
    left_w = usable_w * 0.50
    add_text_box(
        slide, concept_title,
        left=ml, top=text_top + mm(11), width=left_w, height=mm(10),
        font_size=11, bold=True, color=COLOR_BLACK, font_name="Noto Sans JP",
    )
    add_text_box(
        slide, concept_body,
        left=ml, top=text_top + mm(22), width=left_w, height=mm(28),
        font_size=7.5, color=COLOR_BLACK, font_name="Noto Sans JP",
    )

    # 右: 特徴リスト + エリア特性
    right_x = ml + usable_w * 0.52
    right_w = usable_w * 0.48
    feature_text = "\n".join(f"• {f}" for f in features_raw)
    add_text_box(
        slide, feature_text,
        left=right_x, top=text_top + mm(11), width=right_w, height=mm(22),
        font_size=7.5, color=COLOR_GRAY, font_name="Noto Sans JP",
    )
    if area_description:
        add_separator_line(slide, right_x, text_top + mm(34), right_w, COLOR_LIGHT_GRAY)
        add_text_box(
            slide, area_description,
            left=right_x, top=text_top + mm(36), width=right_w, height=mm(18),
            font_size=7, color=COLOR_GRAY, font_name="Noto Sans JP",
        )

    add_header_common(slide, name_en, 3, W, mr)
    return slide


def add_slide4_floorplan(prs, config, text_dir, img_dir, processed_dir):
    """スライド4: FLOOR PLAN（左メイン図面＋右列各階＋下部建物スペック）"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_BG)

    W  = prs.slide_width
    H  = prs.slide_height
    ml = mm(config["slide"]["margin_left_mm"])
    mt = mm(config["slide"]["margin_top_mm"])
    mr = mm(config["slide"]["margin_right_mm"])
    mb = mm(config["slide"]["margin_bottom_mm"])

    name_en       = config["property"]["name_en"]
    property_info = read_md(text_dir / "property_info.md")

    building_type = parse_md_value(property_info, "building_type")
    structure     = parse_md_value(property_info, "structure")
    floors        = parse_md_value(property_info, "floors")
    units         = parse_md_value(property_info, "units")
    building_area = parse_md_value(property_info, "building_area")
    completion    = parse_md_value(property_info, "completion")

    images = find_images(
        REPO_ROOT / config["paths"]["assets"] / "page4_image",
        processed_dir, "page4",
    )

    # ファイル名で断面図と各階平面図を判別
    # 断面図: ファイル名に "section" / "断面" / "cross" を含むもの
    SECTION_KEYWORDS = ("section", "断面", "cross")
    section_img = next(
        (p for p in images if any(kw in p.stem.lower() for kw in SECTION_KEYWORDS)),
        None,
    )
    floor_imgs = [p for p in images if p != section_img]
    # フォールバック: キーワードに一致するファイルがなければ先頭を断面図扱い
    if section_img is None and images:
        section_img = images[0]
        floor_imgs  = images[1:]

    usable_w    = W - ml - mr
    spec_bar_h  = mm(16)
    usable_h    = H - mt - mb - spec_bar_h - mm(4)
    col_gap     = mm(4)

    # 左: 断面図（55%幅）
    left_img_w = usable_w * 0.55
    if section_img:
        add_image_fit(slide, section_img, ml, mt, left_img_w, usable_h)

    # 右列: 各階平面図（最大4枚を縦積み）
    right_x    = ml + left_img_w + col_gap
    right_w    = usable_w - left_img_w - col_gap
    right_imgs = floor_imgs[:4]
    if right_imgs:
        row_gap = mm(3)
        each_h  = (usable_h - row_gap * (len(right_imgs) - 1)) / len(right_imgs)
        for i, img in enumerate(right_imgs):
            img_top = mt + i * (each_h + row_gap)
            add_image_fit(slide, img, right_x, img_top, right_w, each_h)

    # 下部スペックバー
    bar_y = mt + usable_h + mm(4)
    add_separator_line(slide, ml, bar_y, usable_w)

    spec_parts = []
    if building_type: spec_parts.append(building_type)
    if structure:     spec_parts.append(structure)
    if floors:        spec_parts.append(floors)
    if units:         spec_parts.append(units)
    if building_area: spec_parts.append(f"延床 {building_area}")
    if completion:    spec_parts.append(f"竣工 {completion}")

    add_text_box(
        slide, "F L O O R  P L A N",
        left=ml, top=bar_y + mm(2), width=mm(55), height=mm(7),
        font_size=8, bold=True, color=COLOR_BLACK, font_name="Playfair Display",
    )
    if spec_parts:
        add_text_box(
            slide, "  ／  ".join(spec_parts),
            left=ml + mm(57), top=bar_y + mm(3), width=usable_w - mm(57), height=mm(7),
            font_size=6.5, color=COLOR_GRAY, font_name="Noto Sans JP",
        )

    add_header_common(slide, name_en, 4, W, mr)
    return slide


def add_slide5_rentroll(prs, config, text_dir, img_dir, processed_dir):
    """スライド5: RENT ROLL（全幅レイアウト・タイプ別集計＋階ごと部屋グリッド＋合計）"""
    _ = img_dir, processed_dir  # 5Pは画像不使用。インターフェース統一のため引数は保持
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_BG)

    W  = prs.slide_width
    H  = prs.slide_height
    ml = mm(config["slide"]["margin_left_mm"])
    mt = mm(config["slide"]["margin_top_mm"])
    mr = mm(config["slide"]["margin_right_mm"])
    mb = mm(config["slide"]["margin_bottom_mm"])

    name_en         = config["property"]["name_en"]
    rent_roll_md    = read_md(text_dir / "rent_roll.md")
    property_info   = read_md(text_dir / "property_info.md")

    monthly_total   = parse_md_value(rent_roll_md,  "monthly_total")
    annual_total    = parse_md_value(rent_roll_md,  "annual_total")
    annual_income   = parse_md_value(property_info, "annual_income")
    annual_expense  = parse_md_value(property_info, "annual_expense")
    noi             = parse_md_value(property_info, "noi")
    yield_rate      = parse_md_value(property_info, "yield_rate")
    noi_yield       = parse_md_value(property_info, "noi_yield")

    # 部屋テーブル（想定収入表セクションのみ取得）
    room_rows = parse_section_table(rent_roll_md, "想定収入表")

    # room_rows からデータ集計
    grouped_rooms   = group_rooms_by_floor(room_rows)
    type_summary    = derive_type_summary(room_rows)
    contract_status = derive_contract_status(room_rows)

    usable_w = W - ml - mr

    # ── タイトル ──────────────────────────────────────────
    add_text_box(
        slide, "R E N T  R O L L",
        left=ml, top=mt, width=mm(80), height=mm(10),
        font_size=11, bold=True, color=COLOR_BLACK, font_name="Playfair Display",
    )
    add_separator_line(slide, ml, mt + mm(11), usable_w)

    # ── 上段: タイプ別集計（左55%）＋ 契約状況（右45%）────
    upper_top  = mt + mm(14)
    type_w     = usable_w * 0.55
    status_x   = ml + type_w + mm(8)
    status_w   = usable_w - type_w - mm(8)

    # タイプ別集計: C案（テキストボックス＋細線）
    # ヘッダーはゴールドアンダーライン、行は極細グレー線で区切る
    # カラム幅比率: タイプ名(30%) / 間取り(50%) / 戸数(20%)
    row_h     = mm(6)      # 1行あたり高さ
    col_type  = type_w * 0.30
    col_mado  = type_w * 0.50
    col_units = type_w * 0.20

    if type_summary:
        # ヘッダー行
        add_text_box(slide, "TYPE",
                     left=ml, top=upper_top, width=col_type, height=row_h,
                     font_size=5.5, bold=True, color=COLOR_ACCENT,
                     align=PP_ALIGN.LEFT, font_name="Helvetica Neue")
        add_text_box(slide, "LAYOUT",
                     left=ml + col_type, top=upper_top, width=col_mado, height=row_h,
                     font_size=5.5, bold=True, color=COLOR_ACCENT,
                     align=PP_ALIGN.LEFT, font_name="Helvetica Neue")
        add_text_box(slide, "UNITS",
                     left=ml + col_type + col_mado, top=upper_top, width=col_units, height=row_h,
                     font_size=5.5, bold=True, color=COLOR_ACCENT,
                     align=PP_ALIGN.RIGHT, font_name="Helvetica Neue")
        # ヘッダーアンダーライン（ゴールド）
        add_separator_line(slide, ml, upper_top + row_h, type_w, COLOR_ACCENT)

        # データ行
        for i, (type_name, madori, count) in enumerate(type_summary):
            ry = upper_top + row_h + mm(1) + i * row_h
            # 行区切り線（先頭以外・極細グレー）
            if i > 0:
                add_separator_line(slide, ml, ry, type_w, COLOR_ULTRA_LIGHT)
            add_text_box(slide, type_name,
                         left=ml, top=ry, width=col_type, height=row_h,
                         font_size=6, color=COLOR_BLACK,
                         align=PP_ALIGN.LEFT, font_name="Noto Sans JP")
            add_text_box(slide, madori,
                         left=ml + col_type, top=ry, width=col_mado, height=row_h,
                         font_size=6, color=COLOR_GRAY,
                         align=PP_ALIGN.LEFT, font_name="Noto Sans JP")
            add_text_box(slide, str(count),
                         left=ml + col_type + col_mado, top=ry, width=col_units, height=row_h,
                         font_size=6, bold=True, color=COLOR_BLACK,
                         align=PP_ALIGN.RIGHT, font_name="Helvetica Neue")

        # 合計行
        total_count = sum(c for _, _, c in type_summary)
        total_y = upper_top + row_h + mm(1) + len(type_summary) * row_h
        add_separator_line(slide, ml, total_y, type_w, COLOR_LIGHT_GRAY)
        add_text_box(slide, "合計",
                     left=ml, top=total_y + mm(0.5), width=col_type + col_mado, height=row_h,
                     font_size=6, bold=True, color=COLOR_ACCENT,
                     align=PP_ALIGN.LEFT, font_name="Noto Sans JP")
        add_text_box(slide, f"{total_count}戸",
                     left=ml + col_type + col_mado, top=total_y + mm(0.5), width=col_units, height=row_h,
                     font_size=6, bold=True, color=COLOR_ACCENT,
                     align=PP_ALIGN.RIGHT, font_name="Helvetica Neue")

        # upper_h: ヘッダー + データ行 + 合計行 + 余白
        upper_h = row_h + mm(1) + len(type_summary) * row_h + row_h + mm(3)
    else:
        upper_h = mm(25)

    # 契約状況サマリー（ラベル行 + 数値行）
    status_items = [
        ("契約中", contract_status.get("契約中", 0)),
        ("空室",   contract_status.get("空室",   0)),
        ("退去予定", contract_status.get("退去予定", 0)),
        ("申込",   contract_status.get("申込",   0)),
    ]
    s_col_w = status_w / len(status_items)
    for i, (label, value) in enumerate(status_items):
        sx = status_x + i * s_col_w
        add_text_box(
            slide, label,
            left=sx, top=upper_top, width=s_col_w, height=mm(7),
            font_size=6, color=COLOR_GRAY, align=PP_ALIGN.CENTER, font_name="Noto Sans JP",
        )
        val_color = COLOR_ACCENT if label == "空室" and value > 0 else COLOR_BLACK
        add_text_box(
            slide, str(value),
            left=sx, top=upper_top + mm(8), width=s_col_w, height=mm(14),
            font_size=18, bold=True, color=val_color,
            align=PP_ALIGN.CENTER, font_name="Helvetica Neue",
        )

    # ── 財務サマリーボックス（右上エリア・契約状況の下）────────
    # 契約状況の下端: upper_top + mm(8) + mm(14) = upper_top + mm(22)
    fin_top  = upper_top + mm(25)
    fin_x    = status_x
    fin_w    = status_w
    fin_pad  = mm(3)

    # 収支・利回り行の定義（label, value, 大きく表示するか）
    fin_rows = []
    if annual_income:  fin_rows.append(("年間収入",   annual_income,  False))
    if annual_expense: fin_rows.append(("年間支出",   annual_expense, False))
    if noi:            fin_rows.append(("NOI",        noi,            False))
    if yield_rate:     fin_rows.append(("表面利回り",  yield_rate,     True))
    if noi_yield:      fin_rows.append(("NOI利回り",  noi_yield,      True))

    if fin_rows:
        fin_row_h  = mm(8)
        fin_box_h  = fin_pad + len(fin_rows) * fin_row_h + fin_pad
        COLOR_FIN_BG = RGBColor(0xE8, 0xE6, 0xE3)  # 画像参考のグレー

        # 背景矩形
        add_filled_rect(slide, fin_x, fin_top, fin_w, fin_box_h, COLOR_FIN_BG)

        for i, (label, value, large) in enumerate(fin_rows):
            ry = fin_top + fin_pad + i * fin_row_h
            # 行区切り線（先頭以外）
            if i > 0:
                add_separator_line(slide, fin_x + fin_pad, ry, fin_w - fin_pad * 2, COLOR_LIGHT_GRAY)
            add_text_box(
                slide, label,
                left=fin_x + fin_pad, top=ry, width=fin_w * 0.45, height=fin_row_h,
                font_size=5.5, color=COLOR_GRAY,
                align=PP_ALIGN.LEFT, font_name="Noto Sans JP",
            )
            add_text_box(
                slide, value,
                left=fin_x + fin_w * 0.45, top=ry, width=fin_w * 0.55 - fin_pad, height=fin_row_h,
                font_size=7 if large else 5.5,
                bold=large, color=COLOR_BLACK,
                align=PP_ALIGN.RIGHT, font_name="Helvetica Neue",
            )

    # ── 中段: 階ごと部屋グリッド ─────────────────────────
    grid_top  = upper_top + upper_h + mm(5)   # upper_h はテーブル実高さ
    footer_h  = mm(14)
    grid_h    = H - grid_top - footer_h - mb

    floor_label_w  = mm(10)
    grid_content_w = usable_w - floor_label_w
    n_floors       = len(grouped_rooms)

    if n_floors > 0:
        max_rooms   = max(len(rooms) for rooms in grouped_rooms.values())
        cell_w      = grid_content_w / max(max_rooms, 1)
        floor_h     = grid_h / n_floors

        for fi, (floor_no, rooms) in enumerate(grouped_rooms.items()):
            row_top = grid_top + fi * floor_h

            # 行区切り線
            if fi > 0:
                add_separator_line(slide, ml, row_top, usable_w, COLOR_LIGHT_GRAY)

            # 階ラベル
            add_text_box(
                slide, f"{floor_no}F",
                left=ml, top=row_top + floor_h * 0.3,
                width=floor_label_w, height=floor_h * 0.4,
                font_size=7, bold=True, color=COLOR_ACCENT,
                align=PP_ALIGN.CENTER, font_name="Helvetica Neue",
            )

            # 各部屋セル
            # rent_roll.md カラム順: 部屋番号[0] 用途[1] 占有面積[2] タイプ[3] 間取り[4] 賃料[5] 坪単価[6] 状況[7]
            # 表示フォーマット（左列 | 右列）:
            #   部屋番号    | 契約ステータス
            #   タイプ      | 間取り
            #   平米数      | 坪数
            #   賃料        | 賃料/坪
            for ri, room in enumerate(rooms):
                room_no        = room[0] if len(room) > 0 else ""
                area_raw       = room[2] if len(room) > 2 else ""
                type_name      = room[3] if len(room) > 3 else ""
                madori         = room[4] if len(room) > 4 else ""
                rent           = room[5] if len(room) > 5 else ""
                rent_per_tsubo = room[6] if len(room) > 6 else ""
                status         = room[7] if len(room) > 7 else ""

                # 面積を㎡と坪に分割（例: "32.97㎡(9.97坪)" → "32.97㎡" / "9.97坪"）
                area_m = re.match(r"([\d.]+㎡)\(?([\d.]+坪)\)?", area_raw)
                sqm   = area_m.group(1) if area_m else area_raw
                tsubo = area_m.group(2) if area_m else ""

                is_vacant  = (status == "空室")
                main_color = COLOR_ACCENT if is_vacant else COLOR_BLACK

                cell_left = ml + floor_label_w + ri * cell_w
                half_w    = cell_w / 2
                pad       = mm(1.5)  # テキストの左余白

                # 部屋間の縦区切り線（先頭以外）
                if ri > 0:
                    add_vertical_line(
                        slide, cell_left, row_top, floor_h,
                        color=COLOR_LIGHT_GRAY, width_pt=0.5,
                    )

                # 左列・右列の間の内側縦区切り線（さらに薄い）
                add_vertical_line(
                    slide, cell_left + half_w, row_top, floor_h,
                    color=COLOR_ULTRA_LIGHT, width_pt=0.3,
                )

                # 左列: 部屋番号 / タイプ / 平米 / 賃料
                add_multiline_textbox(
                    slide, [
                        (room_no,   5.0, True,  main_color),
                        (type_name, 3.5, False, COLOR_GRAY),
                        (sqm,       3.5, False, COLOR_GRAY),
                        (rent,      4.0, False, main_color),
                    ],
                    left=cell_left + pad, top=row_top,
                    width=half_w - pad, height=floor_h,
                )
                # 右列: ステータス / 間取り / 坪 / 賃料/坪
                add_multiline_textbox(
                    slide, [
                        (status,         5.0, False, main_color),
                        (madori,         3.5, False, COLOR_GRAY),
                        (tsubo,          3.5, False, COLOR_GRAY),
                        (rent_per_tsubo, 4.0, False, COLOR_GRAY),
                    ],
                    left=cell_left + half_w + pad, top=row_top,
                    width=half_w - pad, height=floor_h,
                )

    # ── 下部: 合計 ───────────────────────────────────────
    footer_y = H - footer_h - mb
    add_separator_line(slide, ml, footer_y, usable_w)
    total_text = f"月額総賃料: {monthly_total}"
    if annual_total:
        total_text += f"    年間総賃料: {annual_total}"
    add_text_box(
        slide, total_text,
        left=ml, top=footer_y + mm(2), width=usable_w, height=mm(10),
        font_size=8, bold=True, color=COLOR_BLACK, font_name="Noto Sans JP",
    )

    add_header_common(slide, name_en, 5, W, mr)
    return slide


def add_slide6_location(prs, config, text_dir, img_dir, processed_dir):
    """スライド6: 物件立地（地図2枚縦積み＋物件概要・法令制限・アクセス）"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, COLOR_BG)

    W  = prs.slide_width
    H  = prs.slide_height
    ml = mm(config["slide"]["margin_left_mm"])
    mt = mm(config["slide"]["margin_top_mm"])
    mr = mm(config["slide"]["margin_right_mm"])
    mb = mm(config["slide"]["margin_bottom_mm"])

    name_en       = config["property"]["name_en"]
    property_info = read_md(text_dir / "property_info.md")

    name_ja        = parse_md_value(property_info, "name_ja")           or config["property"]["name_ja"]
    address        = parse_md_value(property_info, "address")
    land_category  = parse_md_value(property_info, "land_category")
    land_rights    = parse_md_value(property_info, "land_rights")
    land_area      = parse_md_value(property_info, "land_area")
    building_type  = parse_md_value(property_info, "building_type")
    structure      = parse_md_value(property_info, "structure")
    floors         = parse_md_value(property_info, "floors")
    units          = parse_md_value(property_info, "units")
    building_area  = parse_md_value(property_info, "building_area")
    completion     = parse_md_value(property_info, "completion")
    zoning         = parse_md_value(property_info, "zoning")
    coverage_ratio = parse_md_value(property_info, "coverage_ratio")
    floor_ratio    = parse_md_value(property_info, "floor_ratio")
    fire_zone      = parse_md_value(property_info, "fire_zone")
    height_zone    = parse_md_value(property_info, "height_zone")
    shadow_reg     = parse_md_value(property_info, "shadow_regulation")
    access_lines   = [
        parse_md_value(property_info, "access1"),
        parse_md_value(property_info, "access2"),
        parse_md_value(property_info, "access3"),
    ]
    access_lines = [a for a in access_lines if a and a != "未記載"]

    page6_dir = REPO_ROOT / config["paths"]["assets"] / "page6_image"
    location_map = find_image_by_prefix(page6_dir, processed_dir, "page6", "location_map")
    site_map     = find_image_by_prefix(page6_dir, processed_dir, "page6", "site_map")

    # 左: 地図エリア（他スライドと共通の余白を使用）
    # location_map（上）＋ site_map（下）の2枚縦積み
    map_left  = ml
    map_top   = mt
    map_w     = W * 0.48 - ml       # 左余白分を除いた幅
    map_h     = H - mt - mb         # 上下余白分を除いた高さ
    map_gap   = mm(3)

    if location_map and site_map:
        each_h = (map_h - map_gap) / 2
        add_image_fit(slide, location_map, map_left, map_top,                map_w, each_h)
        add_image_fit(slide, site_map,     map_left, map_top + each_h + map_gap, map_w, each_h)
    elif location_map:
        add_image_fit(slide, location_map, map_left, map_top, map_w, map_h)
    elif site_map:
        add_image_fit(slide, site_map,     map_left, map_top, map_w, map_h)

    # 右: テキストエリア
    right_x = map_left + map_w + mm(5)
    right_w = W - right_x - mr
    label_w = mm(22)
    value_w = right_w - label_w
    row_h   = mm(7.5)

    # 物件名
    add_text_box(
        slide, name_ja,
        left=right_x, top=mt, width=right_w, height=mm(10),
        font_size=11, bold=True, color=COLOR_BLACK, font_name="Noto Sans JP",
    )
    add_separator_line(slide, right_x, mt + mm(11), right_w)

    cur_y = mt + mm(13)

    def add_info_row(label: str, value: str):
        nonlocal cur_y
        if not value or value == "未記載":
            return
        add_text_box(slide, label,
                     left=right_x, top=cur_y, width=label_w, height=row_h,
                     font_size=6.5, color=COLOR_GRAY, font_name="Noto Sans JP")
        add_text_box(slide, value,
                     left=right_x + label_w, top=cur_y, width=value_w, height=row_h,
                     font_size=6.5, color=COLOR_BLACK, font_name="Noto Sans JP")
        cur_y += row_h

    def add_section_heading(title: str):
        nonlocal cur_y
        add_separator_line(slide, right_x, cur_y + mm(1), right_w, COLOR_LIGHT_GRAY)
        cur_y += mm(4)
        add_text_box(slide, title,
                     left=right_x, top=cur_y, width=right_w, height=mm(5),
                     font_size=6, bold=True, color=COLOR_ACCENT, font_name="Noto Sans JP")
        cur_y += mm(6)

    # 所在地
    add_info_row("所在地", address)

    # 土地情報
    if any([land_category, land_rights, land_area]):
        add_section_heading("土地")
        add_info_row("地目", land_category)
        add_info_row("権利", land_rights)
        add_info_row("面積", land_area)

    # 建物情報
    struct_floors = "  ".join(filter(None, [structure, floors]))
    if any([building_type, struct_floors, units, building_area, completion]):
        add_section_heading("建物")
        add_info_row("種類",   building_type)
        add_info_row("規模・構造", struct_floors)
        add_info_row("戸数",   units)
        add_info_row("延床面積", building_area)
        add_info_row("竣工",   completion)

    # 法令制限
    if any([zoning, coverage_ratio, floor_ratio, fire_zone, height_zone, shadow_reg]):
        add_section_heading("法令制限")
        add_info_row("用途地域", zoning)
        add_info_row("建ぺい率", coverage_ratio)
        add_info_row("容積率",  floor_ratio)
        add_info_row("防火指定", fire_zone)
        add_info_row("高度指定", height_zone)
        add_info_row("日影規制", shadow_reg)

    # 交通アクセス
    if access_lines:
        add_section_heading("交通アクセス")
        for access in access_lines:
            add_text_box(
                slide, f"• {access}",
                left=right_x, top=cur_y, width=right_w, height=row_h,
                font_size=6.5, color=COLOR_BLACK, font_name="Noto Sans JP",
            )
            cur_y += row_h

    add_header_common(slide, name_en, 6, W, mr)
    return slide


# =========================================================
# メイン
# =========================================================

def run():
    config = load_config()
    text_dir      = REPO_ROOT / config["paths"]["intermediates_text"]
    processed_dir = REPO_ROOT / config["paths"]["intermediates_images"]
    output_dir    = REPO_ROOT / config["paths"]["outputs"]
    output_dir.mkdir(parents=True, exist_ok=True)

    print("=== Step 3: PowerPoint生成 ===")

    prs = Presentation()
    prs.slide_width  = Mm(config["slide"]["width_mm"])
    prs.slide_height = Mm(config["slide"]["height_mm"])

    img_dir = REPO_ROOT / config["paths"]["assets"]

    print("スライド1: 表紙")
    add_slide1_cover(prs, config, text_dir, img_dir, processed_dir)
    print("スライド2: 内観")
    add_slide2_interior(prs, config, text_dir, img_dir, processed_dir)
    print("スライド3: コンセプト")
    add_slide3_concept(prs, config, text_dir, img_dir, processed_dir)
    print("スライド4: FLOOR PLAN")
    add_slide4_floorplan(prs, config, text_dir, img_dir, processed_dir)
    print("スライド5: RENT ROLL")
    add_slide5_rentroll(prs, config, text_dir, img_dir, processed_dir)
    print("スライド6: 物件立地")
    add_slide6_location(prs, config, text_dir, img_dir, processed_dir)

    slug     = config["property"]["slug"]
    out_path = output_dir / f"{slug}.pptx"
    prs.save(str(out_path))
    print(f"\nPowerPoint保存: {out_path}")
    print("=== Step 3 完了 ===")
    return out_path


if __name__ == "__main__":
    run()
